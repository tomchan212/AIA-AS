#!/usr/bin/env python3
"""17-page Mandy summary deck, plus a year-by-year appendix."""

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
YEARS = json.loads((ROOT / "src" / "years.json").read_text())
OUT = [
    ROOT / "public" / "mandy.pptx",
    Path("/workspace/blueprint/mandy.pptx"),
]

BG = RGBColor(0xFA, 0xF8, 0xF5)
INK = RGBColor(0x1E, 0x29, 0x3B)
COCOA = RGBColor(0x78, 0x35, 0x0F)
BLUSH = RGBColor(0xFB, 0x71, 0x85)
YELLOW = RGBColor(0xFD, 0xE0, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xFF, 0xE4, 0xE6)


def money(hkd):
    return f"HK$ {hkd:,.2f}"


by_age = {row["age"]: row for row in YEARS}


def at(age):
    return by_age[age]


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, HGT = prs.slide_width, prs.slide_height


def blank():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, HGT)
    fill.line.fill.background()
    fill.fill.solid()
    fill.fill.fore_color.rgb = BG
    return slide


def set_run(paragraph, text, size, bold=False, color=INK, font="Microsoft JhengHei"):
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return run


def textbox(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        set_run(p, line, size, bold, color)
    return box


def card(slide, l, t, w, h, fill=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = INK
    shape.line.width = Pt(2.25)
    return shape


def mandy(slide, l, t, scale=1):
    face = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, Inches(1.15 * scale), Inches(1.15 * scale))
    face.fill.solid()
    face.fill.fore_color.rgb = RGBColor(0xFF, 0xE4, 0xC4)
    face.line.color.rgb = INK
    face.line.width = Pt(2)
    hair = slide.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(0.08 * scale), t - Inches(0.08 * scale), Inches(1.0 * scale), Inches(0.55 * scale))
    hair.fill.solid()
    hair.fill.fore_color.rgb = INK
    hair.line.fill.background()
    for dx in (0.28, 0.72):
        eye = slide.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(dx * scale), t + Inches(0.42 * scale), Inches(0.12 * scale), Inches(0.12 * scale))
        eye.fill.solid()
        eye.fill.fore_color.rgb = INK
        eye.line.fill.background()
    for dx in (0.12, 0.88):
        cheek = slide.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(dx * scale), t + Inches(0.58 * scale), Inches(0.16 * scale), Inches(0.12 * scale))
        cheek.fill.solid()
        cheek.fill.fore_color.rgb = BLUSH
        cheek.line.fill.background()


def footer(slide, n):
    textbox(slide, Inches(0.4), Inches(7.05), Inches(10), Inches(0.3), "說明數字，不是保證，也不是要約。", 11, color=COCOA)
    textbox(slide, Inches(11.4), Inches(7.05), Inches(1.5), Inches(0.3), f"{n} / 17", 12, bold=True, align=PP_ALIGN.RIGHT)


# 1 cover
s = blank()
mandy(s, Inches(5.9), Inches(1.35), 1.6)
textbox(s, Inches(0.6), Inches(3.7), Inches(12), Inches(0.7), "盈御 3 與自願醫保說明", 36, True, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.8), Inches(4.55), Inches(11.7), Inches(0.9), "CHAN MANDY：方案 A 每年儲蓄保費 HK$ 24,000，供 20 年\n方案 B 每年儲蓄保費 HK$ 39,000，供 10 年。兩邊都另加同一列醫療保費。", 18, align=PP_ALIGN.CENTER)
card(s, Inches(3.6), Inches(6.15), Inches(6.1), Inches(0.55), YELLOW)
textbox(s, Inches(3.6), Inches(6.22), Inches(6.1), Inches(0.4), "2019 年 37 歲投保 · 總表由 38 歲起算", 14, True, align=PP_ALIGN.CENTER)
footer(s, 1)

# 2 medical
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "醫療保費會隨年齡上升", 32, True)
textbox(s, Inches(0.5), Inches(1.0), Inches(12), Inches(0.4), "數字是兩份總表裡的「醫療計劃」欄，兩邊相同。", 16, color=COCOA)
points = [(38, at(38)["medical"]), (47, at(47)["medical"]), (60, at(60)["medical"]), (70, at(70)["medical"]), (80, at(80)["medical"]), (100, at(100)["medical"])]
for i, (age, amount) in enumerate(points):
    left = Inches(0.4 + (i % 6) * 2.15)
    card(s, left, Inches(1.7), Inches(2.0), Inches(1.8), WHITE if i % 2 == 0 else PINK)
    textbox(s, left, Inches(1.9), Inches(2.0), Inches(0.4), f"{age} 歲", 16, True, align=PP_ALIGN.CENTER)
    textbox(s, left, Inches(2.45), Inches(2.0), Inches(0.7), money(amount), 14, True, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.5), Inches(4.0), Inches(12), Inches(1.5), "38 歲是 HK$ 6,056.80。100 歲是 HK$ 85,424.80。\n這筆保費每年都要繳，方案 A 和方案 B 用的是同一列。", 22)
footer(s, 2)

# 3 two bills
s = blank()
textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "每年有兩筆保費", 32, True)
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(4.6), WHITE)
textbox(s, Inches(0.8), Inches(1.7), Inches(5.4), Inches(0.5), "自願醫保", 24, True, COCOA)
textbox(s, Inches(0.8), Inches(2.5), Inches(5.4), Inches(2.8), "總表的「醫療計劃」。\n用來付住院和手術的帳單。\n由 38 歲繳到 100 歲，金額逐年增加。", 22)
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(4.6), YELLOW)
textbox(s, Inches(7.1), Inches(1.7), Inches(5.4), Inches(0.5), "盈御 3 儲蓄", 24, True, COCOA)
textbox(s, Inches(7.1), Inches(2.5), Inches(5.4), Inches(2.8), "總表的「盈御多元貨幣計劃 3」。\n方案 A 供 20 年，方案 B 供 10 年。\n供完之後這欄是 0，醫療保費仍然繼續。", 22)
footer(s, 3)

# 4 plan A
s = blank()
textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "方案 A：儲蓄保費供 20 年", 32, True)
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(4.6), WHITE)
textbox(s, Inches(0.8), Inches(1.8), Inches(5.4), Inches(3.5), "38–57 歲，每年 HK$ 24,000\n\n20 年儲蓄本金 HK$ 480,000\n\n58 歲起這欄是 0", 24)
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(4.6), PINK)
textbox(s, Inches(7.1), Inches(1.8), Inches(5.4), Inches(3.6), f"38 歲連醫療一起是\n{money(at(38)['aAnnual'])}\n\n47 歲累計已繳 {money(at(47)['aPaid'])}\n可退保現金價值 {money(at(47)['aSurrender'])}", 20)
footer(s, 4)

# 5 plan B
s = blank()
textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "方案 B：儲蓄保費供 10 年", 32, True)
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(4.6), WHITE)
textbox(s, Inches(0.8), Inches(1.8), Inches(5.4), Inches(3.6), "38–47 歲，每年 HK$ 39,000\n\n10 年儲蓄本金 HK$ 390,000\n\n48 歲起儲蓄保費是 0\n醫療保費繼續", 24)
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(4.6), YELLOW)
textbox(s, Inches(7.1), Inches(1.8), Inches(5.4), Inches(3.6), f"38 歲連醫療一起是\n{money(at(38)['bAnnual'])}\n\n47 歲累計已繳 {money(at(47)['bPaid'])}\n可退保現金價值 {money(at(47)['bSurrender'])}", 20)
footer(s, 5)

# 6 first decade
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "頭 10 年要拿出多少", 32, True)
labels = [("只計醫療", at(38)["medical"], at(47)["medicalCumulative"], WHITE), ("方案 A", at(38)["aAnnual"], at(47)["aPaid"], PINK), ("方案 B", at(38)["bAnnual"], at(47)["bPaid"], YELLOW)]
for i, (name, first, total, fill) in enumerate(labels):
    left = Inches(0.4 + i * 4.25)
    card(s, left, Inches(1.2), Inches(4.0), Inches(3.4), fill)
    textbox(s, left + Inches(0.2), Inches(1.4), Inches(3.6), Inches(0.45), name, 20, True, COCOA)
    textbox(s, left + Inches(0.2), Inches(2.1), Inches(3.6), Inches(1), f"首年\n{money(first)}", 18, True)
    textbox(s, left + Inches(0.2), Inches(3.3), Inches(3.6), Inches(1), f"10 年累計\n{money(total)}", 18, True)
textbox(s, Inches(0.5), Inches(5.0), Inches(12), Inches(1.2), "儲蓄保費每年相差 HK$ 15,000，大約每月 HK$ 1,250。\n醫療保費兩邊一樣，要另計。", 20)
footer(s, 6)

# 7 principal
s = blank()
textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "儲蓄本金", 32, True)
card(s, Inches(0.5), Inches(1.4), Inches(6), Inches(3.2), WHITE)
textbox(s, Inches(0.8), Inches(1.7), Inches(5.4), Inches(0.5), "方案 A", 22, True)
textbox(s, Inches(0.8), Inches(2.5), Inches(5.4), Inches(1.4), "24,000 × 20 年\n= HK$ 480,000", 28, True)
card(s, Inches(6.8), Inches(1.4), Inches(6), Inches(3.2), YELLOW)
textbox(s, Inches(7.1), Inches(1.7), Inches(5.4), Inches(0.5), "方案 B", 22, True)
textbox(s, Inches(7.1), Inches(2.5), Inches(5.4), Inches(1.4), "39,000 × 10 年\n= HK$ 390,000", 28, True)
textbox(s, Inches(0.5), Inches(5.0), Inches(12), Inches(1.3), f"方案 B 的儲蓄本金少 HK$ 90,000。這兩個數未計醫療。\n醫療繳到 100 歲，兩邊都是 {money(at(100)['medicalCumulative'])}。", 20)
footer(s, 7)

# 8 premiums paid
s = blank()
textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.5), "連醫療在內，累計已繳", 32, True)
textbox(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.4), "由 38 歲加到該歲。儲蓄保費停了之後，仍然加上當年的醫療保費。", 16, color=COCOA)
textbox(s, Inches(0.6), Inches(1.45), Inches(2), Inches(0.35), "年齡", 16, True)
textbox(s, Inches(3.0), Inches(1.45), Inches(3), Inches(0.35), "只計醫療", 16, True)
textbox(s, Inches(6.4), Inches(1.45), Inches(3), Inches(0.35), "方案 A", 16, True)
textbox(s, Inches(9.8), Inches(1.45), Inches(3), Inches(0.35), "方案 B", 16, True)
for i, age in enumerate((60, 70, 80, 90, 100)):
    row = at(age)
    top = Inches(1.9 + i * 0.95)
    card(s, Inches(0.4), top, Inches(12.5), Inches(0.85), WHITE if i % 2 == 0 else YELLOW)
    textbox(s, Inches(0.6), top + Inches(0.18), Inches(2), Inches(0.5), f"{age} 歲", 18, True)
    textbox(s, Inches(2.8), top + Inches(0.18), Inches(3.2), Inches(0.5), money(row["medicalCumulative"]), 16)
    textbox(s, Inches(6.2), top + Inches(0.18), Inches(3.2), Inches(0.5), money(row["aPaid"]), 16, True)
    textbox(s, Inches(9.6), top + Inches(0.18), Inches(3.1), Inches(0.5), money(row["bPaid"]), 16, True)
footer(s, 8)

# 9 surrender
s = blank()
textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.5), "可退保現金價值", 32, True)
textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.35), "總表的「可退保現金價值」。當中有非保證部分。", 16, color=COCOA)
for i, age in enumerate((60, 70, 80, 90, 100)):
    row = at(age)
    left = Inches(0.4 + (i % 3) * 4.25)
    top = Inches(1.45 + (i // 3) * 2.5)
    card(s, left, top, Inches(4.05), Inches(2.25), WHITE)
    textbox(s, left + Inches(0.2), top + Inches(0.15), Inches(3.6), Inches(0.4), f"{age} 歲", 18, True)
    textbox(s, left + Inches(0.2), top + Inches(0.65), Inches(3.6), Inches(0.6), f"A  {money(row['aSurrender'])}", 16, True)
    textbox(s, left + Inches(0.2), top + Inches(1.3), Inches(3.6), Inches(0.6), f"B  {money(row['bSurrender'])}", 16, True)
footer(s, 9)

# 10 four ages
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "四個年齡放在一起", 32, True)
textbox(s, Inches(0.6), Inches(1.15), Inches(2.5), Inches(0.4), "年齡", 16, True)
textbox(s, Inches(3.4), Inches(1.15), Inches(4.5), Inches(0.4), "方案 A 退保", 16, True)
textbox(s, Inches(8.2), Inches(1.15), Inches(4.5), Inches(0.4), "方案 B 退保", 16, True)
for i, age in enumerate((47, 57, 60, 100)):
    row = at(age)
    top = Inches(1.7 + i * 1.2)
    card(s, Inches(0.4), top, Inches(12.5), Inches(1.05), WHITE if i % 2 == 0 else PINK)
    textbox(s, Inches(0.6), top + Inches(0.25), Inches(2.4), Inches(0.5), f"{age} 歲", 20, True)
    textbox(s, Inches(3.2), top + Inches(0.25), Inches(4.4), Inches(0.5), money(row["aSurrender"]), 20, True)
    textbox(s, Inches(8.0), top + Inches(0.25), Inches(4.6), Inches(0.5), money(row["bSurrender"]), 20, True)
footer(s, 10)

# 11 47 vs 57
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6), "47 歲仍低於已繳，57 歲高過已繳", 28, True)
card(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(5.2), WHITE)
textbox(s, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.5), "47 歲", 22, True)
textbox(s, Inches(0.6), Inches(2.1), Inches(5.8), Inches(3.6), f"方案 A 已繳 {money(at(47)['aPaid'])}\n退保 {money(at(47)['aSurrender'])}\n\n方案 B 已繳 {money(at(47)['bPaid'])}\n退保 {money(at(47)['bSurrender'])}\n\n兩邊的退保價值都還沒有蓋過已繳保費。", 16)
card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), YELLOW)
textbox(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.5), "57 歲", 22, True)
textbox(s, Inches(7.0), Inches(2.1), Inches(5.7), Inches(3.6), f"方案 A 已繳 {money(at(57)['aPaid'])}\n退保 {money(at(57)['aSurrender'])}\n\n方案 B 已繳 {money(at(57)['bPaid'])}\n退保 {money(at(57)['bSurrender'])}\n\n兩邊的退保價值都高過當時的累計已繳。", 16)
footer(s, 11)

# 12 withdrawable
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "可供提取金額", 32, True)
textbox(s, Inches(0.5), Inches(1.0), Inches(12), Inches(0.7), "兩份總表這一欄的數字相同，由 41 歲列到 79 歲。80 歲起是空白。下面用方案 A 的表。", 16, color=COCOA)
for i, age in enumerate((47, 60, 79)):
    left = Inches(0.5 + i * 4.2)
    card(s, left, Inches(2.0), Inches(3.9), Inches(2.2), WHITE if i != 2 else YELLOW)
    textbox(s, left, Inches(2.2), Inches(3.9), Inches(0.45), f"{age} 歲", 20, True, align=PP_ALIGN.CENTER)
    textbox(s, left, Inches(2.9), Inches(3.9), Inches(0.8), money(at(age)["aWithdrawable"]), 22, True, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.5), Inches(4.6), Inches(12), Inches(1.6), f"70 歲是 {money(at(70)['aWithdrawable'])}。\n這欄不是每年預定領走的金額，總表也沒有列出領走之後還剩多少。", 20)
footer(s, 12)

# 13 age 100
s = blank()
textbox(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "100 歲", 32, True)
card(s, Inches(0.4), Inches(1.3), Inches(6.2), Inches(5), WHITE)
textbox(s, Inches(0.7), Inches(1.5), Inches(5.6), Inches(0.5), "方案 A", 22, True)
textbox(s, Inches(0.7), Inches(2.3), Inches(5.6), Inches(3.4), f"可退保現金價值\n{money(at(100)['aSurrender'])}\n\n連醫療在內已繳\n{money(at(100)['aPaid'])}\n\n當年醫療保費 {money(at(100)['medical'])}", 18)
card(s, Inches(6.8), Inches(1.3), Inches(6.1), Inches(5), YELLOW)
textbox(s, Inches(7.1), Inches(1.5), Inches(5.5), Inches(0.5), "方案 B", 22, True)
textbox(s, Inches(7.1), Inches(2.3), Inches(5.5), Inches(3.4), f"可退保現金價值\n{money(at(100)['bSurrender'])}\n\n連醫療在內已繳\n{money(at(100)['bPaid'])}\n\n當年醫療保費同樣是 {money(at(100)['medical'])}", 18)
footer(s, 13)

# 14 split
s = blank()
textbox(s, Inches(0.5), Inches(0.4), Inches(12), Inches(0.7), "保單可以分拆", 32, True)
textbox(s, Inches(0.5), Inches(1.5), Inches(12), Inches(4.5), "由第 3 個保單年度終結後，或供款期完結後，以較後者為準，可以申請把一部分保單價值分到另一張保單。\n\n分拆要公司核准。分出來的保單，生效日跟原來那張一樣。\n\n這是產品簡介裡的安排，不是這兩份總表已經辦好的事。", 24)
footer(s, 14)

# 15 lock and holiday
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "鎖定，以及延繳保費", 32, True)
card(s, Inches(0.4), Inches(1.15), Inches(6.2), Inches(5.2), WHITE)
textbox(s, Inches(0.6), Inches(1.35), Inches(5.8), Inches(0.5), "紅利鎖定", 22, True)
textbox(s, Inches(0.6), Inches(2.1), Inches(5.8), Inches(3.8), "由第 15 個保單年度終結後起，每個年度終結後 30 日內可鎖定一次。比例 10% 至 70%。\n\n鎖定戶口的利息不是保證。鎖定滿一年可以解鎖，比例 10% 至 100%。", 18)
card(s, Inches(6.8), Inches(1.15), Inches(6.1), Inches(5.2), YELLOW)
textbox(s, Inches(7.0), Inches(1.35), Inches(5.7), Inches(0.5), "延繳保費", 22, True)
textbox(s, Inches(7.0), Inches(2.1), Inches(5.7), Inches(3.8), "只適用 3 年、5 年及 10 年供款，所以只對方案 B。\n\n指定情況包括被裁、結婚、離婚、子女出生或領養、買入住宅、更改主要居住地。每份保單只可申請一次，寬限期最多 365 日。", 18)
footer(s, 15)

# 16 change insured
s = blank()
textbox(s, Inches(0.5), Inches(0.4), Inches(12), Inches(0.7), "可以更換受保人", 32, True)
textbox(s, Inches(0.5), Inches(1.4), Inches(12), Inches(4.8), "由第 1 個保單年度終結後起，受保人在生時，可申請無限次更改受保人，也可以指定第二受保人。要公司批准。\n\n卓越成績獎最高一項，是獲產品簡介所指的全球大學排名首 10 位取錄，獎金 2,800 美元，約 HK$ 21,980（按 7.85）。同一位受保人的所有盈御 3 保單只付一次。\n\n這不是免遺產認證的承諾。", 22)
footer(s, 16)

# 17 close
s = blank()
textbox(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.55), "怎樣看這兩套安排", 32, True)
card(s, Inches(0.4), Inches(1.2), Inches(6.2), Inches(5.2), WHITE)
textbox(s, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.5), "方案 A 較接近的情況", 20, True)
textbox(s, Inches(0.6), Inches(2.2), Inches(5.8), Inches(3.6), f"想每年少付 HK$ 15,000 儲蓄保費。\n可以接受供到 57 歲。\n100 歲的可退保現金價值是\n{money(at(100)['aSurrender'])}。", 20)
card(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.2), YELLOW)
textbox(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.5), "方案 B 較接近的情況", 20, True)
textbox(s, Inches(7.0), Inches(2.2), Inches(5.7), Inches(3.6), f"想 47 歲後不再繳這張儲蓄保單。\n儲蓄本金少 HK$ 90,000。\n100 歲的可退保現金價值是\n{money(at(100)['bSurrender'])}。", 20)
footer(s, 17)


def cell_text(cell, text, size=10, bold=False, align=PP_ALIGN.RIGHT):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    set_run(p, text, size, bold)
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)


chunk = 16
headers = ["年齡", "年份", "醫療", "A 年繳", "A 退保", "B 年繳", "B 退保"]
for start in range(0, len(YEARS), chunk):
    part = YEARS[start:start + chunk]
    s = blank()
    textbox(s, Inches(0.3), Inches(0.12), Inches(12), Inches(0.36), f"附錄 · 每一年（{part[0]['age']}–{part[-1]['age']} 歲）", 20, True)
    textbox(s, Inches(0.3), Inches(0.48), Inches(12.5), Inches(0.28), "年繳 = 當年儲蓄保費 + 醫療保費。可退保現金價值含非保證部分。", 12, color=COCOA)
    table_shape = s.shapes.add_table(len(part) + 1, len(headers), Inches(0.28), Inches(0.85), Inches(12.75), Inches(6.05))
    table = table_shape.table
    widths = [0.9, 1.15, 1.85, 2.05, 2.35, 2.05, 2.35]
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for i, label in enumerate(headers):
        cell = table.cell(0, i)
        cell_text(cell, label, 12, True, PP_ALIGN.CENTER)
        cell.fill.solid()
        cell.fill.fore_color.rgb = YELLOW
    table.rows[0].height = Inches(0.32)
    for r, row in enumerate(part, start=1):
        table.rows[r].height = Inches(0.30)
        values = [
            str(row["age"]),
            str(row["year"]),
            money(row["medical"]),
            money(row["aAnnual"]),
            money(row["aSurrender"]),
            money(row["bAnnual"]),
            money(row["bSurrender"]),
        ]
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell_text(cell, value, 11, False, PP_ALIGN.CENTER if c < 2 else PP_ALIGN.RIGHT)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else PINK

assert at(38)["aAnnual"] == 30056.8
assert at(47)["aPaid"] == 314233.6 and at(47)["bPaid"] == 464233.6
assert at(47)["aSurrender"] == 247604.7 and at(47)["bSurrender"] == 412674.5
assert at(60)["aSurrender"] == 925515 and at(60)["bSurrender"] == 1054160.8
assert at(70)["aSurrender"] == 1422867.45 and at(70)["bSurrender"] == 1610788.6
assert at(100)["aSurrender"] == 6349323.35 and at(100)["bSurrender"] == 7942355.25
assert at(100)["medical"] == 85424.8
assert at(100)["aPaid"] == 2823128.8 and at(100)["bPaid"] == 2733128.8
assert sum(row["aSavings"] for row in YEARS) == 480000
assert sum(row["bSavings"] for row in YEARS) == 390000
assert YEARS[0]["age"] == 38 and YEARS[-1]["age"] == 100 and len(YEARS) == 63

for path in OUT:
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    print("wrote", path, "slides", len(prs.slides))
